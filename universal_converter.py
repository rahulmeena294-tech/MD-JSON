"""
Universal File Converter
Converts Excel (.xlsx, .xls, .csv), Word (.docx), PDF (.pdf), and Text to Markdown (.md) or JSON (.json).
"""

import os
import json
import argparse
from pathlib import Path
from typing import Dict, Any, List, Union
import pandas as pd

# Optional imports handled gracefully
try:
    import docx
except ImportError:
    docx = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None


class UniversalConverter:
    def __init__(self, output_dir: str = None):
        self.output_dir = Path(output_dir) if output_dir else None

    def _get_destination_path(self, input_path: Path, ext: str) -> Path:
        target_dir = self.output_dir if self.output_dir else input_path.parent
        target_dir.mkdir(parents=True, exist_ok=True)
        return target_dir / f"{input_path.stem}_converted.{ext.lstrip('.')}"

    # ---------------- Excel / CSV Handlers ----------------
    def _convert_excel(self, path: Path, output_format: str) -> Path:
        excel_file = pd.ExcelFile(path)
        out_path = self._get_destination_path(path, output_format)

        if output_format == "md":
            md_content = [f"# Data Export: {path.name}\n\n"]
            for sheet_name in excel_file.sheet_names:
                df = excel_file.parse(sheet_name).fillna("")
                md_content.append(f"## Sheet: {sheet_name}\n\n")
                if not df.empty:
                    md_content.append(df.to_markdown(index=False))
                else:
                    md_content.append("*Sheet is empty.*")
                md_content.append("\n\n---\n\n")

            with open(out_path, "w", encoding="utf-8") as f:
                f.writelines(md_content)

        elif output_format == "json":
            sheets_data = {}
            for sheet_name in excel_file.sheet_names:
                df = excel_file.parse(sheet_name)
                sheets_data[sheet_name] = df.to_dict(orient="records")

            with open(out_path, "w", encoding="utf-8") as f:
                json.dump(sheets_data, f, indent=2, default=str)

        return out_path

    def _convert_csv(self, path: Path, output_format: str) -> Path:
        df = pd.read_csv(path).fillna("")
        out_path = self._get_destination_path(path, output_format)

        if output_format == "md":
            md_text = f"# CSV Data: {path.name}\n\n" + df.to_markdown(index=False) + "\n"
            with open(out_path, "w", encoding="utf-8") as f:
                f.write(md_text)
        elif output_format == "json":
            with open(out_path, "w", encoding="utf-8") as f:
                json.dump(df.to_dict(orient="records"), f, indent=2, default=str)

        return out_path

    # ---------------- Word Document Handler ----------------
    def _convert_docx(self, path: Path, output_format: str) -> Path:
        if docx is None:
            raise ImportError("python-docx is not installed. Run 'pip install python-docx'")

        doc = docx.Document(path)
        out_path = self._get_destination_path(path, output_format)

        if output_format == "md":
            md_lines = [f"# Document: {path.name}\n\n"]
            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                # Heading detection based on Word style
                if p.style.name.startswith("Heading 1"):
                    md_lines.append(f"# {text}\n\n")
                elif p.style.name.startswith("Heading 2"):
                    md_lines.append(f"## {text}\n\n")
                elif p.style.name.startswith("Heading 3"):
                    md_lines.append(f"### {text}\n\n")
                else:
                    md_lines.append(f"{text}\n\n")

            # Word tables
            for idx, table in enumerate(doc.tables, start=1):
                data = []
                for row in table.rows:
                    data.append([cell.text.replace("\n", " ").strip() for cell in row.cells])
                if data:
                    df = pd.DataFrame(data[1:], columns=data[0])
                    md_lines.append(f"### Table {idx}\n\n" + df.to_markdown(index=False) + "\n\n")

            with open(out_path, "w", encoding="utf-8") as f:
                f.writelines(md_lines)

        elif output_format == "json":
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            tables_data = []
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    table_rows.append([cell.text.replace("\n", " ").strip() for cell in row.cells])
                if len(table_rows) > 1:
                    tables_data.append(pd.DataFrame(table_rows[1:], columns=table_rows[0]).to_dict(orient="records"))
                elif table_rows:
                    tables_data.append(table_rows)

            payload = {
                "filename": path.name,
                "paragraphs": paragraphs,
                "tables": tables_data
            }
            with open(out_path, "w", encoding="utf-8") as f:
                json.dump(payload, f, indent=2, default=str)

        return out_path

    # ---------------- PDF Handler ----------------
    def _convert_pdf(self, path: Path, output_format: str) -> Path:
        if pdfplumber is None:
            raise ImportError("pdfplumber is not installed. Run 'pip install pdfplumber'")

        out_path = self._get_destination_path(path, output_format)
        pages_content = []

        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables()
                text = page.extract_text() or ""
                pages_content.append({"page": i, "text": text.strip(), "tables": tables})

        if output_format == "md":
            md_lines = [f"# PDF Document: {path.name}\n\n"]
            for item in pages_content:
                md_lines.append(f"## Page {item['page']}\\n\\n")
                if item["text"]:
                    md_lines.append(f"{item['text']}\n\n")
                for t_idx, table in enumerate(item["tables"], start=1):
                    if table and len(table) > 1:
                        # Clean cells
                        cleaned = [[(c.replace("\n", " ").strip() if c else "") for c in row] for row in table]
                        df = pd.DataFrame(cleaned[1:], columns=cleaned[0])
                        md_lines.append(f"**Table {t_idx}**\n\n" + df.to_markdown(index=False) + "\n\n")
                md_lines.append("---\n\n")

            with open(out_path, "w", encoding="utf-8") as f:
                f.writelines(md_lines)

        elif output_format == "json":
            with open(out_path, "w", encoding="utf-8") as f:
                json.dump({"filename": path.name, "pages": pages_content}, f, indent=2, default=str)

        return out_path

    # ---------------- Generic Text / JSON Handlers ----------------
    def _convert_text(self, path: Path, output_format: str) -> Path:
        out_path = self._get_destination_path(path, output_format)
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()

        if output_format == "md":
            with open(out_path, "w", encoding="utf-8") as f:
                f.write(f"# Text File: {path.name}\n\n{content}\n")
        elif output_format == "json":
            with open(out_path, "w", encoding="utf-8") as f:
                json.dump({"filename": path.name, "content": content}, f, indent=2)

        return out_path

    def convert(self, file_path: Union[str, Path], output_format: str = "md") -> Path:
        """
        Converts any supported file to .md or .json.
        """
        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"Input file does not exist: {path}")

        fmt = output_format.lower().strip(".")
        if fmt not in ["md", "json"]:
            raise ValueError("output_format must be either 'md' or 'json'")

        ext = path.suffix.lower()

        if ext in [".xlsx", ".xls"]:
            return self._convert_excel(path, fmt)
        elif ext in [".csv", ".tsv"]:
            return self._convert_csv(path, fmt)
        elif ext == ".docx":
            return self._convert_docx(path, fmt)
        elif ext == ".pdf":
            return self._convert_pdf(path, fmt)
        elif ext in [".txt", ".log"]:
            return self._convert_text(path, fmt)
        elif ext == ".json":
            if fmt == "md":
                with open(path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                out_path = self._get_destination_path(path, "md")
                with open(out_path, "w", encoding="utf-8") as f:
                    f.write(f"# JSON: {path.name}\n\n```json\n{json.dumps(data, indent=2)}\n```\n")
                return out_path
            return path
        else:
            raise NotImplementedError(f"Unsupported file format: {ext}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Universal File to MD / JSON Converter")
    parser.add_argument("input_path", help="Path to the file or directory to convert")
    parser.add_argument("--format", choices=["md", "json"], default="md", help="Target format (md or json)")
    parser.add_argument("--output_dir", default=None, help="Directory where converted files will be saved")

    args = parser.parse_args()
    converter = UniversalConverter(output_dir=args.output_dir)

    target_path = Path(args.input_path)
    if target_path.is_dir():
        supported_exts = {".xlsx", ".xls", ".csv", ".tsv", ".docx", ".pdf", ".txt"}
        files = [p for p in target_path.iterdir() if p.suffix.lower() in supported_exts]
        print(f"Found {len(files)} files in {target_path}...")
        for f in files:
            try:
                res = converter.convert(f, output_format=args.format)
                print(f"[OK] {f.name} -> {res.name}")
            except Exception as e:
                print(f"[ERROR] {f.name}: {e}")
    else:
        result = converter.convert(target_path, output_format=args.format)
        print(f"Converted successfully: {result}")